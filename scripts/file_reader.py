"""Извлечение текста из присланных файлов (Word/Excel/PDF) для разбора задания."""
from __future__ import annotations

from pathlib import Path


def extract_pdf(path: Path) -> str:
    import fitz

    doc = fitz.open(path)
    try:
        return "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()


def extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"[{sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None]
            if values:
                parts.append(" | ".join(values))
    return "\n".join(parts)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix in (".xlsx", ".xlsm"):
        return extract_xlsx(path)
    raise ValueError(f"Формат {suffix} пока не поддерживаю — жду Word (.docx), Excel (.xlsx) или PDF.")
